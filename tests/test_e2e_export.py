import os
import sys
import io

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import openpyxl
from core.data_loader import blocks_to_hourly
from core.sizing import size_pv_system, calc_dimensions, size_battery_bank
from core.dispatch import dispatch_simulation
from core.economics import calc_financial_metrics
from core.exporter import export_to_excel

def test_export_pipeline():
    # 1. Setup mock data
    blocks = [
        (0, 6, 89.0), (6, 9, 98.0), (9, 12, 180.0),
        (12, 15, 208.0), (15, 18, 167.0), (18, 21, 79.0), (21, 24, 59.0)
    ]
    load = blocks_to_hourly(blocks)
    solar = {
        7: 153.33, 8: 296.22, 9: 418.92, 10: 513.06, 11: 572.25,
        12: 592.44, 13: 572.25, 14: 513.06, 15: 418.92, 16: 296.22, 17: 153.33
    }
    modul = {
        "name": "Standard Mono-PERC 550Wp",
        "pmax": 550.0, "voc": 49.5, "vmp": 41.5, "isc": 14.0,
        "kvoc": -0.28, "kvmp": -0.35, "length_mm": 2278, "width_mm": 1134, "cost_per_wp": 3500
    }
    inv = {
        "name": "String Inverter 50kW On-Grid",
        "pac_kw": 50.0, "mppt_min": 500.0, "mppt_max": 850.0, "vdc_max": 1000.0,
        "n_mppt": 4, "imax_mppt": 30.0, "string_per_mppt": 2, "cost_per_w": 1200
    }
    battery_spec = {
        "name": "Lithium LiFePO4 48V 100Ah",
        "voltage_nominal": 48.0,
        "capacity_ah": 100.0,
        "dod_max": 0.85,
        "cost_per_kwh": 3500000
    }

    kwp = 1350.0
    pr = 0.8
    tmin = 18.0
    tmax_sel = 50.0

    pv_res = size_pv_system(kwp, modul, inv, tmin, tmax_sel)
    dims = calc_dimensions(pv_res["n_modul"], modul["length_mm"], modul["width_mm"])
    dims["available_area_m2"] = 10000.0
    
    bat_res = size_battery_bank(daily_kwh=sum(load), autonomy_days=1.0, dod=0.85, v_sys=48.0, battery_spec=battery_spec)
    dispatch_rows, dispatch_summary = dispatch_simulation(
        load, solar, pr, kwp, topology="hybrid", battery_cap_kwh=bat_res["capacity_kwh"]
    )
    econ = calc_financial_metrics(
        kwp=kwp,
        battery_kwh=bat_res["capacity_kwh"],
        annual_energy_kwh=dispatch_summary["total_used"] * 365.0,
        mod_cost_wp=modul["cost_per_wp"],
        inv_cost_w=inv["cost_per_w"],
        battery_cost_kwh=battery_spec["cost_per_kwh"]
    )

    data_bundle = {
        "topology": "Hybrid",
        "kwp": kwp,
        "pr": pr,
        "tmin": tmin,
        "tmax_sel": tmax_sel,
        "psh": dispatch_summary["psh"],
        "module_spec": modul,
        "inverter_spec": inv,
        "battery_spec": battery_spec,
        "battery_res": bat_res,
        "pv_res": pv_res,
        "dims": dims,
        "dispatch_rows": dispatch_rows,
        "dispatch_summary": dispatch_summary,
        "solar_irr": solar,
        "economics": econ
    }

    # 2. Test export to disk
    test_filepath = "test_export_output.xlsx"
    if os.path.exists(test_filepath):
        os.remove(test_filepath)
        
    export_to_excel(data_bundle, test_filepath)
    assert os.path.exists(test_filepath), "File Excel harus terbentuk di disk."
    assert os.path.getsize(test_filepath) > 1000, "Ukuran file Excel harus lebih dari 1KB."

    # 3. Read back with openpyxl and assert sheets and content
    wb = openpyxl.load_workbook(test_filepath)
    sheet_names = wb.sheetnames
    assert "Ringkasan Sizing" in sheet_names
    assert "Dispatch 24 Jam" in sheet_names
    assert "Analisis Finansial" in sheet_names

    ws_dispatch = wb["Dispatch 24 Jam"]
    # Header + 24 rows + total row
    assert ws_dispatch.max_row >= 25

    ws_fin = wb["Analisis Finansial"]
    assert ws_fin.max_row >= 25

    wb.close()
    if os.path.exists(test_filepath):
        os.remove(test_filepath)

    # 4. Test in-memory buffer export (as used in Streamlit download_button)
    buf = io.BytesIO()
    export_to_excel(data_bundle, buf)
    buf.seek(0)
    assert len(buf.getvalue()) > 1000, "Buffer in-memory harus memiliki isi file Excel yang valid."
    wb_mem = openpyxl.load_workbook(buf)
    assert "Ringkasan Sizing" in wb_mem.sheetnames
    wb_mem.close()

if __name__ == "__main__":
    test_export_pipeline()
    print("E2E EXPORT TESTS PASSED")
